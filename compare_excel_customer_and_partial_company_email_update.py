"""
compare_excel_customer_and_partial_company_email_update.py
===========================================================
Vergleicht Kundendaten aus einer Excel-Tabelle mit einem Firmenstamm und
aktualisiert ggf. E-Mail-Adressen.

Dieses Modul enthält keine Rechnungsnummer-Regex und ist vom neuen
PDF-Format nicht direkt betroffen. Es wurde im Rahmen der REPS-Umstellung
geprüft und benötigt keine Anpassung.
"""

import os
import re

import openpyxl
from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Hilfsfunktionen
# ---------------------------------------------------------------------------

def _normalize_company_name(name: str) -> str:
    """Normalisiert einen Firmennamen für den Vergleich.

    Wandelt in Kleinbuchstaben um, entfernt Sonderzeichen und übliche
    Rechtsform-Kürzel (GmbH, AG usw.).
    """
    if not name:
        return ""
    name = name.lower()
    # Rechtsformen entfernen
    name = re.sub(r"\b(gmbh|ag|kg|ohg|gbr|e\.?v\.?|ug|co\.?\s*kg)\b", "", name)
    # Sonderzeichen und mehrfache Leerzeichen
    name = re.sub(r"[^a-zäöüß0-9\s]", " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name


def _emails_match(email1: str, email2: str) -> bool:
    """Prüft ob zwei E-Mail-Adressen (case-insensitive) übereinstimmen."""
    return (email1 or "").strip().lower() == (email2 or "").strip().lower()


# ---------------------------------------------------------------------------
# Vergleichs- und Aktualisierungsfunktionen
# ---------------------------------------------------------------------------

def compare_and_update_emails(
    customer_excel: str,
    company_excel: str,
    output_excel: str | None = None,
    threshold: float = 0.8,
) -> list[dict]:
    """Vergleicht Kundendaten mit Firmenstamm und aktualisiert E-Mails.

    Args:
        customer_excel: Pfad zur Excel-Datei mit Kundendaten.
                        Erwartete Spalten: Name, Email (Rechnungsempfänger).
        company_excel:  Pfad zur Excel-Datei mit Firmenstammdaten.
                        Erwartete Spalten: Firmenname, Email.
        output_excel:   Zielpfad für die aktualisierte Tabelle.
                        Wenn None, wird *customer_excel* überschrieben.
        threshold:      Ähnlichkeitsschwelle (0–1) für den Namensvergleich.

    Returns:
        Liste der aktualisierten Einträge als Dicts.
    """
    if not os.path.isfile(customer_excel):
        raise FileNotFoundError(f"Kunden-Excel nicht gefunden: {customer_excel}")
    if not os.path.isfile(company_excel):
        raise FileNotFoundError(f"Firmen-Excel nicht gefunden: {company_excel}")

    # --- Firmenstamm einlesen ---
    wb_company = load_workbook(company_excel, read_only=True)
    ws_company = wb_company.active
    company_data: list[dict] = []
    headers_company: list[str] = []
    for row_idx, row in enumerate(ws_company.iter_rows(values_only=True)):
        if row_idx == 0:
            headers_company = [str(c).strip() if c else "" for c in row]
            continue
        entry = dict(zip(headers_company, row))
        company_data.append(entry)
    wb_company.close()

    # Firmennamen normalisieren für späteren Vergleich
    for entry in company_data:
        raw = str(entry.get("Firmenname", "") or "")
        entry["_norm"] = _normalize_company_name(raw)

    # --- Kundendaten laden ---
    wb_cust = load_workbook(customer_excel)
    ws_cust = wb_cust.active
    headers_cust: list[str] = []
    updated_rows: list[dict] = []

    for row_idx, row in enumerate(ws_cust.iter_rows()):
        if row_idx == 0:
            headers_cust = [str(c.value).strip() if c.value else "" for c in row]
            continue

        row_data = {headers_cust[i]: row[i].value for i in range(len(headers_cust))}
        customer_name = str(row_data.get("Name", "") or "")
        customer_email = str(row_data.get("Email", "") or "")
        norm_cust = _normalize_company_name(customer_name)

        # Besten Treffer im Firmenstamm suchen
        best_match = _find_best_match(norm_cust, company_data, threshold)

        if best_match:
            company_email = str(best_match.get("Email", "") or "")
            if company_email and not _emails_match(customer_email, company_email):
                # E-Mail aktualisieren
                email_col = headers_cust.index("Email") if "Email" in headers_cust else None
                if email_col is not None:
                    ws_cust.cell(row=row_idx + 1, column=email_col + 1).value = company_email
                row_data["Email"] = company_email
                row_data["_aktualisiert"] = True
                updated_rows.append(row_data)

    # --- Ergebnis speichern ---
    save_path = output_excel or customer_excel
    wb_cust.save(save_path)
    wb_cust.close()

    print(f"{len(updated_rows)} E-Mail(s) aktualisiert. Gespeichert unter: {save_path}")
    return updated_rows


def _find_best_match(
    norm_name: str,
    company_data: list[dict],
    threshold: float,
) -> dict | None:
    """Findet den ähnlichsten Firmeneintrag mittels einfacher Token-Überlappung."""
    if not norm_name:
        return None

    tokens_query = set(norm_name.split())
    if not tokens_query:
        return None

    best_score = 0.0
    best_entry: dict | None = None

    for entry in company_data:
        tokens_company = set((entry.get("_norm") or "").split())
        if not tokens_company:
            continue
        intersection = tokens_query & tokens_company
        union = tokens_query | tokens_company
        score = len(intersection) / len(union) if union else 0.0
        if score > best_score:
            best_score = score
            best_entry = entry

    if best_score >= threshold:
        return best_entry
    return None


# ---------------------------------------------------------------------------
# Direktaufruf
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys

    if len(sys.argv) < 3:
        print(
            "Verwendung: python compare_excel_customer_and_partial_company_email_update.py"
            " <Kunden-Excel> <Firmen-Excel> [Ausgabe-Excel]"
        )
        sys.exit(1)

    kunden_excel = sys.argv[1]
    firmen_excel = sys.argv[2]
    ausgabe_excel = sys.argv[3] if len(sys.argv) > 3 else None

    compare_and_update_emails(kunden_excel, firmen_excel, ausgabe_excel)
