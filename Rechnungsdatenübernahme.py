"""
Rechnungsdatenübernahme.py
==========================
Liest Rechnungs-PDFs aus, extrahiert Rechnungsnummer, Projektname und
Kundendaten und überträgt sie in eine Excel-Datei.

Unterstützte Rechnungsnummern-Formate:
  - Neu:  REPS####  (z. B. REPS1234)  ← bevorzugt
  - Alt:  RE########  (z. B. RE12345678) ← abwärtskompatibel
"""

import os
import re
import glob

import pdfplumber
import openpyxl
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Globale Muster
# ---------------------------------------------------------------------------

# Erkennt beide Formate: REPS1234 (neu) und RE12345678 (alt, abwärtskompatibel)
RE_PATTERN = re.compile(r"\b(?:REPS\d{4}|RE\d{8})\b", re.IGNORECASE)

# Muster für den vollständigen Rechnungsblock in einer Zeile:
#   "Rechnung <Projektname> REPS1234"  oder  "Rechnung <Projektname> RE12345678"
_INVOICE_LINE_PATTERN = re.compile(
    r"Rechnung\s+(.+?)\s+((?:REPS\d{4}|RE\d{8}))\b",
    re.IGNORECASE,
)


# ---------------------------------------------------------------------------
# Hilfsfunktionen
# ---------------------------------------------------------------------------

def _extract_invoice_number(filename: str, full_text: str) -> str:
    """Extrahiert die Rechnungsnummer aus Dateiname oder PDF-Text.

    Vorrang hat der Dateiname, danach der Volltext. Das neue Format REPS####
    wird gegenüber dem alten Format RE######## bevorzugt.
    """
    # 1) Dateiname prüfen – bevorzugt neues Format
    m_file_new = re.search(r"\bREPS\d{4}\b", filename, flags=re.IGNORECASE)
    if m_file_new:
        return m_file_new.group(0).upper()

    m_file_old = re.search(r"\bRE\d{8}\b", filename, flags=re.IGNORECASE)
    if m_file_old:
        return m_file_old.group(0).upper()

    # 2) Volltext prüfen – bevorzugt neues Format
    m_text_new = re.search(r"\bREPS\d{4}\b", full_text or "", flags=re.IGNORECASE)
    if m_text_new:
        return m_text_new.group(0).upper()

    m_text_old = re.search(r"\bRE\d{8}\b", full_text or "", flags=re.IGNORECASE)
    if m_text_old:
        return m_text_old.group(0).upper()

    return ""


def _extract_project_and_invoice_from_line(full_text: str) -> tuple[str, str]:
    """Parst eine Zeile der Form 'Rechnung <Projektname> REPS1234'.

    Gibt (projektname, rechnungsnummer) zurück oder ("", "").
    """
    m = _INVOICE_LINE_PATTERN.search(full_text)
    if m:
        projekt = m.group(1).strip()
        rechnung_nr = m.group(2).upper()
        return projekt, rechnung_nr
    return "", ""


def _extract_customer_from_text(full_text: str) -> str:
    """Versucht, den Kundennamen aus dem PDF-Text zu extrahieren.

    Sucht nach typischen Feldern wie 'Kunde:', 'Auftraggeber:' oder
    einem Adressblock nach dem Rechnungsblock.
    """
    for label in ("Kunde", "Auftraggeber", "Rechnungsempfänger", "An"):
        m = re.search(
            rf"{label}\s*[:\-]?\s*(.+?)(?:\n|$)",
            full_text,
            flags=re.IGNORECASE,
        )
        if m:
            candidate = m.group(1).strip()
            if candidate:
                return candidate

    # Fallback: erste nicht-leere Zeile nach dem Rechnungsblock
    m_block = _INVOICE_LINE_PATTERN.search(full_text)
    if m_block:
        rest = full_text[m_block.end():]
        for line in rest.splitlines():
            line = line.strip()
            if line and not re.match(r"^(Datum|Seite|MwSt|Gesamt)", line, re.IGNORECASE):
                return line

    return ""


# ---------------------------------------------------------------------------
# Haupt-Extraktionsfunktion
# ---------------------------------------------------------------------------

def extract_invoice_data_to_excel(pdf_path: str, excel_path: str) -> dict:
    """Liest Rechnungsdaten aus *pdf_path* und schreibt sie in *excel_path*.

    Gibt ein dict mit den extrahierten Feldern zurück.
    """
    if not os.path.isfile(pdf_path):
        raise FileNotFoundError(f"PDF nicht gefunden: {pdf_path}")

    # --- PDF-Text extrahieren ---
    full_text = ""
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            full_text += page_text + "\n"

    filename = os.path.basename(pdf_path)

    # --- Rechnungsnummer & Projekt ---
    projekt, rechnung_nr = _extract_project_and_invoice_from_line(full_text)

    # Falls die Inline-Erkennung leer ist, Fallback auf allgemeine Muster
    if not rechnung_nr:
        rechnung_nr = _extract_invoice_number(filename, full_text)

    if not projekt:
        # Projekt aus Dateinamen ableiten (Format: Rechnung_REPS1234_Projektname.pdf)
        m_proj = re.search(
            r"Rechnung_(?:REPS\d{4}|RE\d{8})_(.+)\.pdf",
            filename,
            flags=re.IGNORECASE,
        )
        if m_proj:
            projekt = m_proj.group(1).replace("_", " ").strip()

    # --- Kunde ---
    kunde = _extract_customer_from_text(full_text)

    # --- Datum ---
    datum = ""
    m_datum = re.search(
        r"(?:Rechnungsdatum|Datum)\s*[:\-]?\s*(\d{1,2}[./]\d{1,2}[./]\d{2,4})",
        full_text,
        flags=re.IGNORECASE,
    )
    if m_datum:
        datum = m_datum.group(1).strip()

    # --- Betrag (Gesamtbetrag brutto) ---
    betrag = ""
    m_betrag = re.search(
        r"(?:Gesamtbetrag|Rechnungsbetrag|Betrag)\s*[:\-]?\s*([\d.,]+)\s*€?",
        full_text,
        flags=re.IGNORECASE,
    )
    if m_betrag:
        betrag = m_betrag.group(1).strip()

    # --- Projekt aus Text (zusätzlich/alternativ) ---
    if not projekt:
        projekt_match = re.search(
            r"Rechnung\s+(.*?(?:REPS\d{4}|RE\d{8})\b)",
            full_text,
            flags=re.IGNORECASE | re.DOTALL,
        )
        if projekt_match:
            # Alles zwischen "Rechnung" und der Rechnungsnummer ist der Projektname
            raw = projekt_match.group(1)
            nr_in_raw = re.search(r"\b(?:REPS\d{4}|RE\d{8})\b", raw, re.IGNORECASE)
            if nr_in_raw:
                projekt = raw[: nr_in_raw.start()].strip()

    data = {
        "Rechnungsnummer": rechnung_nr,
        "Projekt": projekt,
        "Kunde": kunde,
        "Datum": datum,
        "Betrag": betrag,
        "Quelle": filename,
    }

    # --- Excel schreiben / aktualisieren ---
    _write_to_excel(data, excel_path)

    return data


def _write_to_excel(data: dict, excel_path: str) -> None:
    """Schreibt *data* als neue Zeile in die Excel-Tabelle *excel_path*."""
    headers = ["Rechnungsnummer", "Projekt", "Kunde", "Datum", "Betrag", "Quelle"]

    if os.path.isfile(excel_path):
        wb = load_workbook(excel_path)
        ws = wb.active
        # Prüfen ob Header vorhanden, sonst einfügen
        if ws.max_row == 0 or ws.cell(row=1, column=1).value != "Rechnungsnummer":
            ws.insert_rows(1)
            for col, h in enumerate(headers, start=1):
                ws.cell(row=1, column=col, value=h)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rechnungen"
        ws.append(headers)

    # Neue Datenzeile anhängen
    ws.append([data.get(h, "") for h in headers])
    wb.save(excel_path)


# ---------------------------------------------------------------------------
# Batch-Verarbeitung
# ---------------------------------------------------------------------------

def process_folder(pdf_folder: str, excel_path: str) -> list[dict]:
    """Verarbeitet alle PDFs in *pdf_folder* und überträgt die Daten nach *excel_path*."""
    pdf_files = glob.glob(os.path.join(pdf_folder, "*.pdf"))
    results = []
    for pdf_file in sorted(pdf_files):
        try:
            data = extract_invoice_data_to_excel(pdf_file, excel_path)
            results.append(data)
            print(f"[OK] {os.path.basename(pdf_file)} → {data['Rechnungsnummer']}")
        except Exception as exc:  # noqa: BLE001
            print(f"[FEHLER] {os.path.basename(pdf_file)}: {exc}")
    return results


# ---------------------------------------------------------------------------
# Direktaufruf
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys

    if len(sys.argv) < 3:
        print("Verwendung: python Rechnungsdatenübernahme.py <PDF-Ordner> <Excel-Datei>")
        sys.exit(1)

    folder = sys.argv[1]
    excel = sys.argv[2]
    ergebnisse = process_folder(folder, excel)
    print(f"\n{len(ergebnisse)} Rechnung(en) verarbeitet.")
